"""The plugins: schedule and reminders, the knowledge library, documents, backup, coach, web and music."""
import base64
import json
import zipfile
from datetime import date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from airtac import agent, briefing, config, memory, tools
from airtac.main import app
from airtac.scheduler import Scheduler

client = TestClient(app, base_url="http://localhost")


def plugin_module(name):
    import importlib
    return importlib.import_module(f"plugins.{name}")


# ---------------------------------------------------------------- schedule

def test_briefing_reads_your_day():
    store = plugin_module("schedule").store
    today = date.today()
    store.add_event("DBMS lab record", (today - timedelta(days=2)).isoformat())
    store.add_event("Mock interview", today.isoformat(), "14:00")
    store.add_event("Submit form", (today + timedelta(days=1)).isoformat())
    b = briefing.build(datetime.combine(today, datetime.min.time()).replace(hour=9))
    assert b["text"].startswith("Good morning, Ali.")
    assert "1 item is overdue: DBMS lab record." in b["text"]
    assert "Today at 14:00: Mock interview, in 5 hours." in b["text"] and "Today at 2 pm: Mock interview" in b["spoken"]
    assert "It's 9 am on " in b["spoken"]
    assert "Tomorrow: Submit form." in b["text"]


def test_schedule_editor_and_plans():
    r = client.post("/api/plugins/schedule", json={"title": "OS exam", "date": "2027-03-10", "category": "exam"})
    event_id = r.json()["id"]
    assert client.get("/api/plugins/schedule").json()["events"][0]["title"] == "OS exam"
    plans = {t["name"] for t in client.get("/api/plugins/schedule/templates").json()}
    assert {"internship-2027", "hackathon-2026"} <= plans
    added = client.post("/api/plugins/schedule/template", json={"name": "internship-2027"}).json()["added"]
    assert added > 5 and client.post("/api/plugins/schedule/template", json={"name": "internship-2027"}).json()["added"] == 0
    assert client.post(f"/api/plugins/schedule/{event_id}", json={"done": True}).json() == {"ok": True}
    assert client.delete(f"/api/plugins/schedule/{event_id}").json() == {"ok": True}
    assert client.post("/api/plugins/schedule", json={"title": "x", "date": "2027-02-30"}).status_code == 422


def test_reminders_go_off_once_and_the_morning_briefing_too():
    sched = plugin_module("schedule")
    now = datetime(2026, 10, 12, 7, 40)
    sched.store.add_event("Team call", "2026-10-12", "07:50")
    sched.store.add_event("Drink water", "2026-10-12", "07:41", "reminder")
    s = Scheduler()
    made = s.tick(now, clock=0)
    titles = [n["title"] for n in memory.notifications_after(0)]
    assert "Team call at 07:50" in titles and "Your day" in titles      # 10 min before; after 07:30
    assert "Reminder: Drink water" in titles                            # within its minute
    assert made == 3 and s.tick(now + timedelta(minutes=1), clock=120) == 0  # nothing twice


def test_a_spoken_reminder_with_a_time_lands_on_the_schedule():
    reply = agent.chat("Remind me to pay the exam fee tomorrow at 6 pm")
    assert reply.steps[0].tool == "reminder" and "6:00 pm" in reply.spoken
    ev = plugin_module("schedule").store.list_events(category="reminder")[0]
    assert ev["time"] == "18:00" and ev["title"].startswith("Pay the exam fee")


# ---------------------------------------------------------------- knowledge

def test_library_indexes_and_answers_with_sources():
    k = plugin_module("knowledge")
    folder = config.KNOWLEDGE_DIR / "coding"
    folder.mkdir(parents=True)
    (folder / "big-o.md").write_text("# Big O\n\nBinary search runs in O(log n) time on a sorted array, because each "
                                     "step halves the range.\n\nA hash table gives O(1) average lookups.\n", encoding="utf-8")
    (config.KNOWLEDGE_DIR / "cardio.txt").write_text("The cardiac cycle has two phases: systole, when the ventricles "
                                                     "contract, and diastole, when they relax and fill.", encoding="utf-8")
    result = k.index_all()
    assert result["added"] == 2 and k.store.counts()["docs"] == 2
    found = k.search("how fast is binary search on a sorted array")
    assert found[0]["title"] == "big o" and found[0]["shelf"] == "coding"
    ctx = k.passages("What are the phases of the cardiac cycle, systole and diastole?")
    assert ctx and "systole" in ctx["text"] and ctx["sources"][0]["title"] == "cardio"
    assert k.index_all()["added"] == 0  # nothing changed
    (config.KNOWLEDGE_DIR / "cardio.txt").unlink()
    assert k.index_all()["removed"] == 1


def test_library_passages_reach_the_answer():
    k = plugin_module("knowledge")
    config.KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
    (config.KNOWLEDGE_DIR / "ohm.md").write_text("Ohm's law: the voltage across a resistor equals current times "
                                                 "resistance, V = I R.", encoding="utf-8")
    k.index_all()

    class Recorder:
        def __init__(self):
            self.sent = None

        def chat(self, messages, tools, max_tokens=600):
            from airtac.llm import LLMReply
            self.sent = messages[0]["content"]
            return LLMReply(text="V = I R.")
    llm = Recorder()
    reply = agent.chat("What does Ohm's law say about voltage and resistance?", llm=llm)
    assert "V = I R" in llm.sent and reply.sources[0]["title"] == "ohm"


def test_meaning_search_with_an_embedding_model(monkeypatch):
    k = plugin_module("knowledge")
    vocab = ["heart", "attack", "myocardial", "infarction", "python", "list"]

    def fake_embed(texts):  # "heart attack" and "myocardial infarction" point the same way
        out = []
        for t in texts:
            low = t.lower()
            out.append([float(("heart" in low or "myocardial" in low)), float(("attack" in low or "infarction" in low)),
                        0.0, 0.0, float("python" in low), float("list" in low)])
        return out
    monkeypatch.setattr(k, "embed", fake_embed)
    config.KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
    (config.KNOWLEDGE_DIR / "mi.txt").write_text("Myocardial infarction needs urgent care: call an ambulance.", encoding="utf-8")
    (config.KNOWLEDGE_DIR / "py.txt").write_text("A Python list keeps items in order and can grow.", encoding="utf-8")
    k.index_all()
    found = k.search("signs of a heart attack", 1)
    assert found[0]["title"] == "mi" and found[0]["cosine"] > 0.9
    assert vocab  # (the fake model's vocabulary, for the reader)


def test_upload_to_the_library():
    data = base64.b64encode(b"Dijkstra finds shortest paths from one source when no edge is negative.").decode()
    r = client.post("/api/plugins/knowledge/upload", json={"name": "graphs.txt", "shelf": "Coding", "data_base64": data})
    assert r.json()["path"] == "coding/graphs.txt" and r.json()["passages"] == 1
    assert client.post("/api/plugins/knowledge/upload", json={"name": "x.exe", "data_base64": data}).status_code == 422
    hits = client.get("/api/plugins/knowledge/search", params={"q": "shortest paths"}).json()["results"]
    assert hits and "path" not in hits[0]


def test_long_text_is_cut_into_overlapping_passages():
    k = plugin_module("knowledge")
    text = " ".join(f"Sentence number {i} talks about a topic in some detail." for i in range(200))
    parts = k.split([("p. 3", text)])
    assert len(parts) > 5 and all(len(t) <= k.CHUNK + 5 for _, t in parts) and parts[0][0] == "p. 3"


# ---------------------------------------------------------------- documents

def test_word_slides_and_excel_open_again():
    import docx
    from openpyxl import load_workbook
    from pptx import Presentation
    tools.run_tool("make_document", {"title": "Plan", "body": "## Week 1\n- Arrays\nRevise daily."})
    tools.run_tool("make_slides", {"title": "Two pointers", "slides": "## Idea\n- Two indexes\n## Example\n- Pair sum"})
    tools.run_tool("make_spreadsheet", {"title": "Budget", "rows": "Item,Amount\nRent,5000\nFood,3000"})
    out = config.OUTPUTS_DIR / "documents"
    d = docx.Document(str(next(out.glob("plan-*.docx"))))
    assert any(par.text == "Week 1" for par in d.paragraphs)
    deck = Presentation(str(next(out.glob("two-pointers-*.pptx"))))
    assert len(deck.slides) == 3 and deck.slides[1].shapes.title.text == "Idea"
    ws = load_workbook(next(out.glob("budget-*.xlsx"))).active
    assert ws["A1"].value == "Item" and ws["B2"].value == 5000 and ws["A1"].font.bold


# ---------------------------------------------------------------- backup, coach

def test_backup_has_the_database_and_never_the_keys(tmp_path):
    memory.save_fact("my exam is on 12 Nov")
    r = tools.run_tool("backup_now", {})
    zips = list(config.BACKUP_DIR.glob("airtac-backup-*.zip"))
    assert r.ok and len(zips) == 1
    names = zipfile.ZipFile(zips[0]).namelist()
    assert "data/airtac.db" in names and not any(n.endswith(".env") for n in names)
    upload = tools.run_tool("backup_upload", {})
    assert not upload.ok and "remote" in upload.spoken


def test_weekly_review_and_training_export():
    first = agent.chat("Remember my exam is on 12 Nov")
    memory.add_feedback(first.message_id, True)
    r = tools.run_tool("weekly_review", {})
    assert "rated 1 answers: 1 good" in r.spoken
    out = tools.run_tool("export_training_data", {})
    assert out.ok
    line = next((config.OUTPUTS_DIR / "coach").glob("training-*.jsonl")).read_text(encoding="utf-8").splitlines()[0]
    assert json.loads(line)["messages"][0]["content"] == "Remember my exam is on 12 Nov"


# ---------------------------------------------------------------- web and music

def test_web_search_and_reader_offline():
    data = client.get("/api/plugins/web/search", params={"q": "nemotron"}).json()
    assert data["sample"] and len(data["results"]) == 3
    page = client.get("/api/plugins/web/read", params={"url": "https://nebius.com/"}).json()
    assert "Token Factory" in page["text"]
    assert client.get("/api/plugins/web/read", params={"url": "https://example.com/"}).status_code == 422


def test_web_providers_follow_the_settings(monkeypatch):
    web = plugin_module("web")
    monkeypatch.setitem(config.RAW.setdefault("plugins", {}), "web", {"searxng_url": "http://localhost:8888"})
    monkeypatch.setenv("TAVILY_API_KEY", "x")
    assert web.providers() == ["searxng", "tavily", "ddg", "wikipedia"]


def test_music_plays_from_your_folder_only(tmp_path, monkeypatch):
    music = tmp_path / "Music" / "Lofi"
    music.mkdir(parents=True)
    (music / "rainy lofi beat.mp3").write_bytes(b"ID3fake")
    monkeypatch.setitem(config.RAW.setdefault("plugins", {}), "music", {"folder": str(tmp_path / "Music")})
    plugin_module("music")._index.clear()
    r = tools.run_tool("play_music", {"query": "lofi"})
    url = r.data["audio"]
    assert r.ok and "rainy lofi beat" in r.spoken
    assert client.get(url).content == b"ID3fake"
    assert client.get("/api/plugins/music/play/0000000000000000").status_code == 404


def test_desktop_opens_only_what_you_allowed(monkeypatch):
    desk = plugin_module("desktop")
    opened = []
    monkeypatch.setattr(desk, "launch", opened.append)
    assert desk.open_app("notepad").ok and opened == ["notepad.exe"]
    assert not desk.open_app("regedit").ok and opened == ["notepad.exe"]


def test_mcp_never_offers_write_or_device_tools():
    from airtac import mcp_server
    names = {t.name for t in mcp_server.exposed_tools()}
    assert "web_search" in names and not names & {"create_issue", "open_app", "open_website", "backup_upload"}


# ---------------------------------------------------------------- EduSkills

def test_eduskills_from_enrolment_to_report():
    import docx
    edu = plugin_module("eduskills")
    start = (date.today() + timedelta(days=3)).isoformat()
    end = (date.today() + timedelta(days=40)).isoformat()
    r = tools.run_tool("eduskills_add", {"domain": "cybersecurity", "partner": "Palo Alto", "start": start, "end": end,
                                         "assessment": end, "modules_total": 8})
    assert r.ok and "on your schedule" in r.spoken
    titles = [e["title"] for e in plugin_module("schedule").store.list_events(category="internship")]
    assert "EduSkills Cybersecurity starts" in titles and "EduSkills Cybersecurity final assessment" in titles
    u = tools.run_tool("eduskills_update", {"domain": "cyber", "modules_done": 5})
    assert "3 modules to go" in u.spoken
    tools.run_tool("eduskills_log", {"text": "Finished the firewall basics module"})
    s = tools.run_tool("eduskills_status", {})
    assert "1 EduSkills internship in progress" in s.spoken
    rep = tools.run_tool("eduskills_report", {"domain": "cybersecurity"})
    assert rep.ok and "fill in" in rep.spoken
    doc = docx.Document(str(next((config.OUTPUTS_DIR / "documents").glob("cybersecurity-internship-report-*.docx"))))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Finished the firewall basics module" in text and "Palo Alto" in text
    assert edu.store.find("palo")["modules_done"] == 5


def test_eduskills_by_voice_and_in_the_briefing():
    reply = agent.chat("Add my EduSkills AI-ML internship starting tomorrow")
    assert reply.steps[0].tool == "eduskills_add"
    assert agent.chat("How are my EduSkills internships going?").steps[0].tool == "eduskills_status"
    assert agent.chat("EduSkills steps").steps[0].tool == "eduskills_guide"
    edu = plugin_module("eduskills")
    i = edu.store.all_internships()[0]
    edu.store.update(i["id"], assessment=(date.today() + timedelta(days=2)).isoformat())
    b = briefing.build()
    assert "EduSkills Ai-ml assessment" in b["text"]
    page = client.get("/api/apps/eduskills").json()
    assert page["stats"][0]["value"] == 1
