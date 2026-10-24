"""Part C: attachments (checks, readers, index, chat context) and the knowledge search with citations."""
import base64
import io
import zipfile

import pytest
from fastapi.testclient import TestClient

from airtac import agent, config
from airtac.main import app
from plugins import attachments as att
from plugins import knowledge

here = lambda: TestClient(app, base_url="http://localhost")   # noqa: E731


def make_pdf(pages: list[str]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_docx(text: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def make_pptx(slides: list[str]) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for t in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(section, name, raw):
    r = here().post("/api/upload", json={"section": section, "name": name, "data_base64": base64.b64encode(raw).decode()})
    return r


def read(r):
    assert r.status_code == 200, r.text
    return att.read_now(r.json()["id"])


# ---------------------------------------------------------------- checks

@pytest.mark.parametrize("name,raw,why", [
    ("virus.exe", b"MZ\x90\x00", "takes"),
    ("fake.pdf", b"MZ\x90\x00 not a pdf", "real PDF"),
    ("photo.png", b"GIF89a....", "real PNG"),
    ("notes.txt", b"hello\x00\x00world", "real TXT"),
    ("empty.md", b"", "empty"),
])
def test_bad_files_are_refused(name, raw, why):
    with pytest.raises(att.Refused, match=why):
        att.check(name, raw)


def test_a_zip_bomb_disguised_as_docx_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("word/document.xml", b"0" * 30_000_000)   # 30 MB of zeros packs to ~30 KB
    with pytest.raises(att.Refused, match="zip bomb"):
        att.check("report.docx", out.getvalue())


def test_file_names_are_made_safe():
    assert att.safe_name("..\\..\\Windows\\evil name?.PDF") == "evil-name.pdf"
    assert att.safe_name("../../.env") == "env"
    assert "/" not in att.safe_name("a/b/c.txt")


def test_sections_and_sizes_are_checked(monkeypatch):
    assert upload("../etc", "a.txt", b"hi").status_code == 422
    monkeypatch.setattr(att, "MAX_MB", 0.000001)
    assert upload("chat", "a.txt", b"hello there").status_code == 422


# ---------------------------------------------------------------- readers

def test_pdf_pages_are_read_indexed_and_kept_in_the_section_folder():
    r = read(upload("schedule", "dbms notes.pdf", make_pdf(["Normal forms: 1NF removes repeating groups.",
                                                            "BCNF: every determinant is a candidate key."])))
    assert r["status"] == "ready" and r["parts"] == 2 and "BCNF" in att.get(r["id"])["text"]
    assert (config.DB_PATH.parent / "uploads" / "schedule").glob(f"{r['id']}-dbms-notes.pdf")
    assert [d for d in knowledge.store.docs(attached=True) if d["shelf"] == "@schedule"]
    assert not [d for d in knowledge.store.docs() if d["shelf"].startswith("@")]      # not mixed into the library


def test_word_powerpoint_excel_csv_and_text_are_read():
    for name, raw, word in [("a.docx", make_docx("Dijkstra uses a priority queue"), "priority"),
                            ("b.pptx", make_pptx(["AVL rotations", "Left-right case"]), "Left-right"),
                            ("c.xlsx", make_xlsx([["Subject", "Exam"], ["DBMS", "12 Nov"]]), "DBMS | 12 Nov"),
                            ("d.csv", b"topic,week\nheaps,11\n", "heaps | 11"),
                            ("e.md", "# Héllo\nunicode ok".encode(), "Héllo")]:
        r = read(upload("study", name, raw))
        assert r["status"] == "ready", (name, r["message"])
        assert word in att.get(r["id"])["text"], name


def test_old_office_files_say_how_to_fix_them():
    r = read(upload("study", "old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600))
    assert r["status"] == "error" and "save it as .docx" in r["message"]


def test_photos_use_ocr_and_say_so_when_it_is_missing(monkeypatch):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(buf, "PNG")
    monkeypatch.setattr(att, "ocr_bytes", lambda data: "DBMS mid exam 12 Nov 2026")
    r = read(upload("schedule", "timetable.png", buf.getvalue()))
    assert r["status"] == "ready" and "12 Nov" in att.get(r["id"])["text"]

    def missing(data):
        raise att.Unavailable("OCR isn't installed")
    monkeypatch.setattr(att, "ocr_bytes", missing)
    r = read(upload("schedule", "timetable2.png", buf.getvalue()))
    assert r["status"] == "error" and "OCR isn't installed" in r["message"]
    assert here().get(f"/api/plugins/attachments/{r['id']}/file").headers["content-type"] == "image/png"


def test_audio_is_transcribed_or_explained(monkeypatch):
    wav = b"RIFF\x24\x00\x00\x00WAVEfmt " + b"\0" * 40
    monkeypatch.setattr(att, "transcribe", lambda path: [("0:00", "Lecture on AVL trees and rotations")])
    r = read(upload("study", "lecture.wav", wav))
    assert r["status"] == "ready" and "AVL" in att.get(r["id"])["text"]


def test_ocr_really_reads_a_picture_when_rapidocr_is_installed():
    pytest.importorskip("rapidocr")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (700, 120), "white")
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 40)
    except OSError:
        pytest.skip("no font for the test picture")
    ImageDraw.Draw(img).text((10, 30), "AVL tree insert", fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    assert "AVL" in att.ocr_bytes(buf.getvalue())


# ---------------------------------------------------------------- the chat answers from attachments

def test_attached_text_goes_into_the_models_instructions():
    r = read(upload("chat", "plan.txt", b"Week 1 starts on 19 Oct: arrays and prefix sums."))
    messages, sources = agent._build_context("s", "When does week 1 start?", "general", [r["id"]])
    assert "19 Oct" in messages[0]["content"] and "plan.txt" in messages[0]["content"]
    assert sources[0]["title"] == "plan.txt"


def test_long_attachments_send_the_matching_pages_only(monkeypatch):
    pages = [f"Filler page {i}. " + "Lorem ipsum dolor sit amet. " * 30 for i in range(12)]
    pages[8] = "AVL tree insertion: insert like a BST, then rebalance with rotations (LL, RR, LR, RL)."
    r = read(upload("research", "dsa.pdf", make_pdf(pages)))
    found = att.context_for([r["id"]], "how to insert into an AVL tree rotations")
    assert "rebalance with rotations" in found["text"] and any(s["where"] == "p. 9" for s in found["sources"])
    assert len(found["text"]) < 8000


def test_unready_or_failed_files_are_named_not_hidden():
    r = upload("chat", "old.ppt", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600).json()
    att.read_now(r["id"])
    assert "couldn't be read" in att.context_for([r["id"]], "x")["text"]


def test_chat_route_accepts_attachments():
    r = read(upload("chat", "a.txt", b"My roll number is 23091A3301."))
    res = here().post("/api/chat", json={"message": "what is my roll number", "attachments": [r["id"]]})
    assert res.status_code == 200


# ---------------------------------------------------------------- knowledge.search with citations

def test_knowledge_search_returns_file_page_and_score(tmp_path):
    path = knowledge.library() / "coding" / "trees.pdf"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(make_pdf(["Binary heaps.", "An AVL tree keeps balance factor -1, 0 or 1 after every insert."]))
    knowledge.index_file(path)
    found = knowledge.search("AVL balance factor", 3, collection="coding")
    assert found and found[0]["file"] == "trees.pdf" and found[0]["page"] == 2 and found[0]["score"] > 0
    answer = here().get("/api/plugins/knowledge/search", params={"q": "AVL balance factor"}).json()["answer"]
    assert "[1]" in answer["text"] and answer["sources"][0]["page"] == 2
    assert knowledge.search("AVL", 3, collection="@nothing") == []


# ---------------------------------------------------------------- a GitHub repository as Research's sub-brain

def fake_github(monkeypatch, files: dict, meta: dict | None = None):
    """GitHub's repo info, its list of files, and the raw files, without the internet."""
    import httpx
    info = {"full_name": "mytechnotalent/Reverse-Engineering", "default_branch": "main", "size": 351123, "private": False,
            "license": {"spdx_id": "Apache-2.0"}, "html_url": "https://github.com/mytechnotalent/Reverse-Engineering",
            "description": "A free reverse engineering tutorial", **(meta or {})}
    tree = {"tree": [{"path": n, "type": "blob", "size": len(d)} for n, d in files.items()] + [{"path": "0x01_hello_world", "type": "tree"}]}
    calls = []

    def respond(url, **k):
        calls.append(url)
        req = httpx.Request("GET", url)
        if "/git/trees/" in url:
            return httpx.Response(200, json=tree, request=req)
        if url.startswith("https://api.github.com/repos/"):
            return httpx.Response(200, json=info, request=req)
        path = url.split("/main/", 1)[1]
        data = files.get(path.replace("%20", " "))
        return httpx.Response(200 if data is not None else 404, content=data.encode() if isinstance(data, str) else (data or b""), request=req)

    class Client:
        def __init__(self, **k):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get(self, url, **k):
            return respond(url, **k)
    monkeypatch.setattr(knowledge.httpx, "get", respond)
    monkeypatch.setattr(knowledge.httpx, "Client", Client)
    return calls


def test_a_repository_becomes_a_collection_with_its_licence_and_no_programs(monkeypatch):
    calls = fake_github(monkeypatch, {
        "README.md": "# Reverse Engineering\nA free tutorial for x86, ARM and RISC-V.",
        "0x01_hello_world/README.md": "The stack frame: push rbp, mov rbp, rsp sets up the frame pointer.",
        "tool.exe": b"MZ\x90\x00",
        "../evil.md": "climbing out",
        ".github/secret.md": "hidden",
    })
    out = knowledge.add_repo("https://github.com/mytechnotalent/Reverse-Engineering", "reverse-engineering")
    assert out["collection"] == "reverse-engineering" and out["files"] == 2 and out["licence"] == "Apache-2.0"
    folder = knowledge.library() / "reverse-engineering" / "mytechnotalent-reverse-engineering"
    assert (folder / "SOURCE.md").read_text().count("Apache-2.0") == 1
    assert not list(folder.rglob("*.exe")) and not (knowledge.library() / "evil.md").exists() and not list(folder.rglob("secret.md"))
    assert calls[0] == "https://api.github.com/repos/mytechnotalent/Reverse-Engineering"
    assert not any(c.endswith("tool.exe") for c in calls)                 # programs aren't even downloaded
    assert {"id": "reverse-engineering", "docs": 3} in knowledge.collections()
    hits = knowledge.search("stack frame rbp", 3, collection="reverse-engineering")
    assert hits and "push rbp" in hits[0]["text"]


def test_bad_repository_names_and_private_repos_are_refused(monkeypatch):
    for bad in ("not a repo", "https://evil.example/a/b", "a/b/c"):
        with pytest.raises(ValueError):
            knowledge.repo_name(bad)
    fake_github(monkeypatch, {"README.md": "x"}, {"private": True})
    with pytest.raises(ValueError, match="public"):
        knowledge.add_repo("someone/private-thing")


def test_research_answers_from_the_chosen_sub_brain(monkeypatch):
    from plugins import research
    fake_github(monkeypatch, {"README.md": "Ghidra decompiles a binary into C-like pseudocode for analysis."})
    knowledge.add_repo("mytechnotalent/Reverse-Engineering", "reverse-engineering")
    text, used = research.gather("what does ghidra do with a binary", None, "reverse-engineering")
    assert "pseudocode" in text and used == ["reverse-engineering collection"]
    r = here().post("/api/plugins/research/copilot", json={"action": "ask", "question": "what does ghidra do with a binary",
                                                           "collection": "reverse-engineering"}).json()
    assert "pseudocode" in r["markdown"]
    assert here().post("/api/plugins/research/copilot", json={"action": "ask", "question": "x", "collection": "../x"}).status_code == 422


def test_research_ask_is_not_limited_to_notes_and_knows_about_you():
    from airtac import memory
    from plugins import research
    memory.save_profile({"name": "Ali", "studies": "B.Tech CSE (AI & ML), 2nd year", "college": "RGM, Nandyal"})
    text, used = research.gather("who is ali", None)
    assert "RGM, Nandyal" in text and "About you" in used
    assert "About the user" in agent._build_context("s", "who am I?", "general")[0][0]["content"]
    r = here().post("/api/plugins/research/copilot", json={"action": "ask", "question": "who is ali", "text": ""}).json()
    assert "RGM" in r["markdown"] or "don't" not in r["markdown"]


def test_github_route_is_for_this_pc_only():
    assert TestClient(app).post("/api/plugins/knowledge/github", json={"repo": "a/b"}).status_code in (401, 403)
    assert here().post("/api/plugins/knowledge/github", json={"repo": "not valid!"}).status_code == 422
