"""Documents: Word files, PowerPoint slides and Excel sheets, made on this laptop.

Write the content in plain text and AIRTAC lays it out:
  ## A heading (or a new slide)      - a bullet point      anything else is a paragraph
Sheets take rows as lines with commas (the first line is the header) or as a JSON list of lists.
Uses python-docx, python-pptx and openpyxl (all MIT). Files go to outputs/documents/.
"""
import json
import re
from datetime import datetime

from airtac.plugin_api import Tool, ToolResult, capitalize, plugin, schema

p = plugin(__name__)


def _slug(title: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:40] or "document"


def _path(title: str, ext: str):
    return p.output_dir() / f"{_slug(title)}-{datetime.now():%Y%m%d-%H%M%S}.{ext}"


def _blocks(body: str) -> list[tuple[str, str]]:
    """[(kind, text)] with kind heading, bullet or text."""
    out = []
    for line in (body or "").splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            out.append(("heading", s.lstrip("#").strip()))
        elif re.match(r"^[-*•]\s+", s):
            out.append(("bullet", re.sub(r"^[-*•]\s+", "", s)))
        else:
            out.append(("text", s))
    return out


def make_document(title: str, body: str) -> ToolResult:
    import docx
    title = (title or "Document").strip()[:120]
    doc = docx.Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Made by AIRTAC Advance on {datetime.now():%d %b %Y}").italic = True
    for kind, text in _blocks(body):
        if kind == "heading":
            doc.add_heading(text, level=1)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    path = _path(title, "docx")
    doc.save(path)
    return ToolResult(f"I made a Word document called {title}. It's in your files.",
                      f"Saved {path.name}\nOpen: /files/documents/{path.name}")


def make_slides(title: str, slides: str) -> ToolResult:
    from pptx import Presentation
    from pptx.util import Pt
    title = (title or "Slides").strip()[:120]
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = f"Made by AIRTAC Advance, {datetime.now():%d %b %Y}"
    current = None
    count = 0
    for kind, text in _blocks(slides):
        if kind == "heading" or current is None:
            current = deck.slides.add_slide(deck.slide_layouts[1])
            current.shapes.title.text = text if kind == "heading" else title
            frame = current.placeholders[1].text_frame
            frame.clear()
            count += 1
            if kind == "heading":
                continue
        para = frame.add_paragraph() if frame.paragraphs[0].text else frame.paragraphs[0]
        para.text = text
        para.font.size = Pt(20)
    path = _path(title, "pptx")
    deck.save(path)
    return ToolResult(f"I made a deck called {title} with {count + 1} slides. It's in your files.",
                      f"Saved {path.name}\nOpen: /files/documents/{path.name}")


def _rows(rows: str) -> list[list]:
    rows = (rows or "").strip()
    if rows.startswith("["):
        try:
            data = json.loads(rows)
            return [list(r) for r in data if isinstance(r, (list, tuple))]
        except ValueError:
            pass
    out = []
    for line in rows.splitlines():
        if line.strip():
            cells = [c.strip() for c in (line.split("\t") if "\t" in line else line.split(","))]
            out.append([float(c) if re.fullmatch(r"-?\d+(\.\d+)?", c) else c for c in cells])
    return out


def make_spreadsheet(title: str, rows: str) -> ToolResult:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    title = (title or "Sheet").strip()[:120]
    data = _rows(rows)
    if not data:
        return ToolResult("What should go in the sheet? Give me rows, one per line, with commas between cells.")
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Sheet1"
    for r in data:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    path = _path(title, "xlsx")
    wb.save(path)
    return ToolResult(f"I made a spreadsheet called {title} with {len(data) - 1} rows. It's in your files.",
                      f"Saved {path.name}\nOpen: /files/documents/{path.name}")


p.add(Tool("make_document", "Create a Word (.docx) document from a title and plain-text body ('## ' headings, '- ' bullets).",
           schema({"title": {"type": "string"}, "body": {"type": "string"}}, required=("title", "body")),
           "local", make_document))
p.add(Tool("make_slides", "Create a PowerPoint (.pptx) deck: each '## ' line starts a slide, '- ' lines are its bullets.",
           schema({"title": {"type": "string"}, "slides": {"type": "string"}}, required=("title", "slides")),
           "local", make_slides))
p.add(Tool("make_spreadsheet", "Create an Excel (.xlsx) sheet from rows (one per line, cells separated by commas; first row is the header).",
           schema({"title": {"type": "string"}, "rows": {"type": "string"}}, required=("title", "rows")),
           "local", make_spreadsheet))


@p.dashboard
def page(pg: dict) -> None:
    files = sorted(p.output_dir().glob("*.*"), key=lambda f: f.stat().st_mtime, reverse=True)
    kinds = {"docx": "Word", "pptx": "Slides", "xlsx": "Excel"}
    pg["stats"] = [{"label": label, "value": sum(1 for f in files if f.suffix == f".{ext}")} for ext, label in kinds.items()]
    pg["lists"] = [{"title": "Latest", "empty": "Nothing yet. Ask: make slides about two pointers.",
                    "items": [{"title": f.name, "meta": kinds.get(f.suffix[1:], ""), "url": f"/files/documents/{f.name}"}
                              for f in files[:10]]}]


# ---------------------------------------------------------------- rules for the no-model mode
# Without a model AIRTAC can't write the content, so these make a labelled starter you can edit.

def _topic(text: str, word: str) -> str:
    m = re.search(rf"\b{word}\w*\s+(?:about|on|for|of)\s+(.+)$", text, re.I)
    return capitalize(m.group(1).strip(" .?!")) if m else "Notes"


def _doc_args(text, _m):
    t = _topic(text, "(?:document|doc|word)")
    return {"title": t, "body": f"## {t}\nWrite your notes here.\n## Key points\n- First point\n- Second point\n"
                                "(Made without a model: start Ollama and ask again for real content.)"}


def _slides_args(text, _m):
    t = _topic(text, "(?:slides?|deck|presentation|ppt)")
    return {"title": t, "slides": f"## What is {t}?\n- One line definition\n## How it works\n- Step one\n- Step two\n"
                                  "## Summary\n- (Made without a model: start Ollama for real content.)"}


def _sheet_args(text, _m):
    t = _topic(text, "(?:spreadsheet|sheet|excel|table)")
    return {"title": t, "rows": "Item,Amount,Notes\nExample,0,Edit me"}


p.mock_rule(7.2, ("word document", "docx", "make a document", "write a document"), "make_document", _doc_args)
p.mock_rule(7.3, ("slides", "presentation", "ppt", "pptx", "deck"), "make_slides", _slides_args)
p.mock_rule(7.4, ("spreadsheet", "excel", "xlsx"), "make_spreadsheet", _sheet_args)
